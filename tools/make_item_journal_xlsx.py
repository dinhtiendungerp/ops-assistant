"""Sinh file Excel Item Journal de import vao Business Central, tao ra Item Ledger Entry
co Lot No. va Expiration Date that, dung theo kich ban da thong nhat trong make_fixtures.py.

Vi sao lam duong nay: khong can build extension AL, khong can compiler, khong can quyen dac biet.
Chi can import config package vao Item Journal roi Post.

Nguyen tac de so lieu khop voi fixtures:
  - Moi thang sinh mot lo san xuat cho moi mat hang. Lo do nhap vao tung dia diem dung bang
    luong ban cua chinh thang do tai dia diem do, nen ban het, khong con ton du.
  - Ton kho cuoi ky la cac lo "hien tai" lay nguyen tu gen_stock, nhap gan day, giu nguyen
    han dung da ghim de cac tinh huong demo (can han, qua han, ton thua) khong bi troi.
  - Moi dong ban deu ghi ro Lot No. cua lo thang do, nen khong phu thuoc thu tu FEFO cua BC.

Xuat mot file cho moi thang de post theo dung thu tu thoi gian, khong dung mot batch 9000 dong.

Chay:
    python tools/make_item_journal_xlsx.py --out /mnt/user-data/outputs/marou-data
"""
from __future__ import annotations

import argparse
import random
import sys
import unicodedata
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python"))

from bc_agent.make_fixtures import ITEMS, SEED, STORES, TODAY, WH, gen_sales, gen_stock  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

COLUMNS = [
    "Batch Name", "Posting Date", "Entry Type", "Document No.", "Item No.", "Description",
    "Variant Code", "Location Code", "Bin Code", "Quantity", "Unit of Measure Code",
    "Unit Amount", "Amount", "Discount Amount", "Unit Cost", "Gen. Bus. Posting Group",
    "Applies-to Entry", "Serial No.", "Lot No.", "Expiration Date", "Warranty Date",
    "Department Code", "Project Code", "Customergroup Code", "Area Code",
    "Businessgroup Code", "Salescampaign Code", "ItemDescription",
]

BATCH = "NWVDEMO"
UOM = "PCS"
MARKUP = 2.2          # gia ban le tren gia von, chi de Value Entry co doanh thu
HISTORY_DAYS = 180

ITEM_BY_NO = {i[0]: i for i in ITEMS}


def ascii_slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.upper() if ch.isalnum())


def lot_code(item_no: str, d: date) -> str:
    """Mot lo san xuat cho moi mat hang moi thang, dung chung cho moi dia diem.
    Nho vay demo truy xuat nguon goc chay duoc: mot lo di den nhieu cua hang."""
    short = item_no.replace("MAR-", "").replace("-", "")[:8]
    return f"L{d:%y%m}-{short}"


def month_key(d: date) -> tuple[int, int]:
    return d.year, d.month


def month_start(y: int, m: int, floor: date) -> date:
    return max(date(y, m, 1), floor)


def build_rows() -> list[dict]:
    rng = random.Random(SEED)
    sales = gen_sales(rng, days=HISTORY_DAYS)
    lots = gen_stock(rng)
    start = TODAY - timedelta(days=HISTORY_DAYS)

    # ---- gom luong ban theo thang x item x dia diem de biet can nhap bao nhieu
    need: dict[tuple[tuple[int, int], str, str], int] = defaultdict(int)
    for r in sales:
        d = date.fromisoformat(r["date"])
        need[(month_key(d), r["item_no"], r["location_code"])] += r["qty"]

    rows: list[dict] = []

    # ---- nhap kho dau moi thang, dung bang luong se ban trong thang do
    for (mk, item_no, loc), qty in sorted(need.items()):
        if qty <= 0:
            continue
        recv = month_start(mk[0], mk[1], start)
        _, desc, _, cost, _, shelf = ITEM_BY_NO[item_no]
        rows.append(make_row(
            posting_date=recv, entry_type="Purchase", item_no=item_no, description=desc,
            location=loc, qty=qty, unit_amount=cost, unit_cost=cost,
            lot_no=lot_code(item_no, recv), expiration=recv + timedelta(days=shelf),
        ))

    # ---- ban tung ngay, ghi ro lo cua thang do
    for r in sorted(sales, key=lambda x: (x["date"], x["item_no"], x["location_code"])):
        d = date.fromisoformat(r["date"])
        item_no = r["item_no"]
        _, desc, _, cost, _, shelf = ITEM_BY_NO[item_no]
        recv = month_start(d.year, d.month, start)
        rows.append(make_row(
            posting_date=d, entry_type="Sale", item_no=item_no, description=desc,
            location=r["location_code"], qty=r["qty"],
            unit_amount=round(cost * MARKUP, 0), unit_cost=cost,
            lot_no=lot_code(item_no, recv), expiration=recv + timedelta(days=shelf),
        ))

    # ---- ton kho hien tai: cac lo ghim san, nhap gan day, giu nguyen han dung
    for lot in sorted(lots, key=lambda x: (x["item_no"], x["location_code"], x["lot_no"])):
        if lot["qty"] <= 0:
            continue
        item_no = lot["item_no"]
        _, desc, _, cost, _, shelf = ITEM_BY_NO[item_no]
        recv = lot["expiration_date"] - timedelta(days=shelf)
        recv = min(max(recv, start), TODAY)
        rows.append(make_row(
            posting_date=recv, entry_type="Purchase", item_no=item_no, description=desc,
            location=lot["location_code"], qty=lot["qty"], unit_amount=cost, unit_cost=cost,
            lot_no=lot["lot_no"], expiration=lot["expiration_date"],
        ))

    rows.sort(key=lambda r: (r["Posting Date"], 0 if r["Entry Type"] == "Purchase" else 1,
                             r["Item No."], r["Location Code"]))
    return rows


def make_row(posting_date: date, entry_type: str, item_no: str, description: str, location: str,
             qty: int, unit_amount: float, unit_cost: float, lot_no: str, expiration: date) -> dict:
    return {
        "Batch Name": BATCH,
        "Posting Date": posting_date,
        "Entry Type": entry_type,
        "Document No.": f"MJ{posting_date:%y%m%d}",
        "Item No.": item_no,
        "Description": description,
        "Variant Code": "",
        "Location Code": location,
        "Bin Code": "",
        "Quantity": qty,
        "Unit of Measure Code": UOM,
        "Unit Amount": unit_amount,
        "Amount": round(unit_amount * qty, 0),
        "Discount Amount": 0,
        "Unit Cost": unit_cost,
        "Gen. Bus. Posting Group": "",
        "Applies-to Entry": "",
        "Serial No.": "",
        "Lot No.": lot_no,
        "Expiration Date": expiration,
        "Warranty Date": "",
        "Department Code": "",
        "Project Code": "",
        "Customergroup Code": "",
        "Area Code": "",
        "Businessgroup Code": "",
        "Salescampaign Code": "",
        "ItemDescription": description,
    }


HEAD_FILL = PatternFill("solid", fgColor="1F3864")


def write_sheet(ws, rows: list[dict]) -> None:
    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
    for c, name in enumerate(COLUMNS, start=1):
        width = 14
        if name in ("Description", "ItemDescription"):
            width = 24
        if name in ("Item No.", "Lot No.", "Location Code"):
            width = 16
        ws.column_dimensions[get_column_letter(c)].width = width
    for col in ("B", "T", "U"):
        for cell in ws[col][1:]:
            cell.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    rows = build_rows()
    by_month: dict[tuple[int, int], list[dict]] = defaultdict(list)
    for r in rows:
        by_month[month_key(r["Posting Date"])].append(r)

    summary = []
    for idx, mk in enumerate(sorted(by_month), start=1):
        chunk = by_month[mk]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Item Journals"
        ws.sheet_properties.tabColor = "1F3864"
        write_sheet(ws, chunk)
        name = f"{idx:02d}-ItemJournals-{mk[0]}-{mk[1]:02d}.xlsx"
        wb.save(out / name)
        n_recv = sum(1 for r in chunk if r["Entry Type"] == "Purchase")
        summary.append((name, len(chunk), n_recv, len(chunk) - n_recv))

    print(f"{len(rows)} dong, {len(by_month)} file")
    for name, total, recv, sale in summary:
        print(f"  {name:38s} {total:6d} dong  nhap {recv:5d}  ban {sale:6d}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
