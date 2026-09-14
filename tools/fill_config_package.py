"""Do du lieu demo Marou vao Config Package MD_v1.1.xlsx.

Sua bon sheet: 27 Item, 5700 Stockkeeping Unit, 6502 Item Tracking Code, 83 Item Journal Line.
Cac sheet khac giu nguyen byte, khong dung toi.

Khong dung openpyxl de ghi, vi no lam mat xmlMaps va tableParts nen BC khong nhan file nua.
Ghi thang vao XML cua tung sheet, giu nguyen dong 1 (ma package, ten bang, table id) va dong 3
(dong tieu de), chi dung lai vung du lieu tu dong 4.
"""
from __future__ import annotations

import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

# Chay: python tools/fill_config_package.py <package BC xuat ra.xlsx> <file ket qua.xlsx>
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("MD_v1.1.xlsx")
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("MD_v1.2-marou.xlsx")
DEMO = Path(__file__).resolve().parents[1] / "demo-data-nwv"

TEMPLATE, BATCH = "ITEM", "NWVDEMO"

ROW_RE = re.compile(r"<x:row r=\"(\d+)\"[^>]*>(.*?)</x:row>", re.S)
CELL_RE = re.compile(
    r"<x:c r=\"([A-Z]+)\d+\"[^>]*>(?:<x:is><x:t[^>]*>(.*?)</x:t></x:is>)?</x:c>", re.S
)
REF_RE = re.compile(r"ref=\"A3:([A-Z]+)\d+\"")


def col_letter(i: int) -> str:
    s = ""
    while i > 0:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(letter: str) -> int:
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n


def esc(v: str) -> str:
    return v.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def unesc(v: str) -> str:
    return v.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")


def cell(ref: str, value: str) -> str:
    return ('<x:c r="' + ref + '" s="1" t="inlineStr"><x:is>'
            '<x:t xml:space="preserve">' + esc(value) + "</x:t></x:is></x:c>")


def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else ("%.5f" % v).rstrip("0").rstrip(".")
    return str(v)


def parse_sheet(xml: str):
    i = xml.index("<x:sheetData>") + len("<x:sheetData>")
    j = xml.index("</x:sheetData>")
    head, body, tail = xml[:i], xml[i:j], xml[j:]
    rows = {int(m.group(1)): m.group(0) for m in ROW_RE.finditer(body)}
    header = {}
    for letter, text in CELL_RE.findall(rows[3]):
        name = unesc(text or "").strip()
        if name:
            header[name] = col_index(letter)
    data = []
    for n in sorted(rows):
        if n >= 4:
            data.append({col_index(l): unesc(t or "") for l, t in CELL_RE.findall(rows[n])})
    return head, tail, rows[1], rows[3], header, data


def build_sheet(head, tail, row1, row3, rows) -> str:
    out = [head, row1, row3]
    for n, r in enumerate(rows, start=4):
        cells = "".join(cell(col_letter(c) + str(n), v) for c, v in sorted(r.items()))
        out.append('<x:row r="' + str(n) + '">' + cells + "</x:row>")
    out.append(tail)
    return "".join(out)


def retable(xml: str, last_row: int) -> str:
    return REF_RE.sub(lambda m: 'ref="A3:' + m.group(1) + str(last_row) + '"', xml)


def read_master():
    wb = openpyxl.load_workbook(DEMO / "00-MasterData.xlsx", data_only=True)

    def sheet(name):
        rows = list(wb[name].iter_rows(values_only=True))
        hdr = list(rows[0])
        return [dict(zip(hdr, r)) for r in rows[1:] if r[0]]

    return sheet("Item"), sheet("Stockkeeping Unit")


def read_journals():
    out = []
    for f in sorted(DEMO.glob("0[1-7]-ItemJournals-*.xlsx")):
        wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows))
        n = 0
        for r in rows:
            if r[0] is None:
                continue
            out.append(dict(zip(hdr, r)))
            n += 1
        print("  " + f.name + ": " + str(n) + " dong")
        wb.close()
    return out


def main():
    items, skus = read_master()
    print("master data: %d item, %d stockkeeping unit" % (len(items), len(skus)))
    print("item journal:")
    journals = read_journals()
    print("tong: %d dong" % len(journals))
    # Giu nguyen thu tu cua bay file nguon, khong sap xep lai. Thu tu do la thu tu vat ly:
    # nhap vao kho truoc, xuat kho sau, nhan hang tai cua hang, roi moi ban. Sap xep lai theo
    # ten Entry Type se dao "Negative Adjmt." len truoc "Purchase" va lam am ton trong ngay.
    dates = [r["Posting Date"] for r in journals]
    assert all(dates[i] <= dates[i + 1] for i in range(len(dates) - 1)), "ngay khong tang dan"

    zin = zipfile.ZipFile(SRC)
    order = zin.namelist()
    parts = {n: zin.read(n) for n in order}
    zin.close()

    def load(n):
        return parts[n].decode("utf-8-sig")

    report = {}

    # 6502 Item Tracking Code: chi giu LOTALLEXP, bat co tao Lot No. Information khi post
    head, tail, r1, r3, hdr, data = parse_sheet(load("xl/worksheets/sheet16.xml"))
    c_code = hdr["Code"]
    c_flag = hdr["Create Lot No. Info. on posting"]
    keep = [r for r in data if r.get(c_code) == "LOTALLEXP"]
    assert len(keep) == 1, "tim thay %d dong LOTALLEXP" % len(keep)
    keep[0][c_flag] = "true"
    parts["xl/worksheets/sheet16.xml"] = build_sheet(head, tail, r1, r3, keep).encode("utf-8")
    parts["xl/tables/table16.xml"] = retable(load("xl/tables/table16.xml"), 3 + len(keep)).encode("utf-8")
    report["6502 Item Tracking Code"] = len(keep)

    # 27 Item: giu dung 21 ma cua bo demo, cap nhat cac cot can sua
    head, tail, r1, r3, hdr, data = parse_sheet(load("xl/worksheets/sheet4.xml"))
    by_no = {str(r["No."]): r for r in items}
    c_no = hdr["No."]
    cols = ["Unit Cost", "Last Direct Cost", "Unit Price", "Item Tracking Code",
            "Lot Nos.", "Replenishment System", "Reordering Policy"]
    for name in cols:
        assert name in hdr, "sheet Item khong co cot " + name
    keep = []
    for r in data:
        src = by_no.get(r.get(c_no))
        if not src:
            continue
        for name in cols:
            if src.get(name) is not None:
                r[hdr[name]] = fmt(src[name])
        keep.append(r)
    thieu = sorted(set(by_no) - set(r[c_no] for r in keep))
    assert not thieu, "khong thay ma trong package: %s" % thieu
    parts["xl/worksheets/sheet4.xml"] = build_sheet(head, tail, r1, r3, keep).encode("utf-8")
    parts["xl/tables/table4.xml"] = retable(load("xl/tables/table4.xml"), 3 + len(keep)).encode("utf-8")
    report["27 Item"] = len(keep)

    # 5700 Stockkeeping Unit: thay bang cac dong cua bo demo
    head, tail, r1, r3, hdr, data = parse_sheet(load("xl/worksheets/sheet15.xml"))
    sku_cols = ["Location Code", "Item No.", "Variant Code", "Replenishment System",
                "Reordering Policy", "Reorder Point", "Reorder Quantity",
                "Safety Stock Quantity", "Maximum Inventory", "Transfer-from Code"]
    for name in sku_cols:
        assert name in hdr, "sheet Stockkeeping Unit khong co cot " + name
    rows = []
    for s in skus:
        rows.append({hdr[name]: fmt(s.get(name)) for name in sku_cols})
    parts["xl/worksheets/sheet15.xml"] = build_sheet(head, tail, r1, r3, rows).encode("utf-8")
    parts["xl/tables/table15.xml"] = retable(load("xl/tables/table15.xml"), 3 + len(rows)).encode("utf-8")
    report["5700 Stockkeeping Unit"] = len(rows)

    # 83 Item Journal Line
    # "Expiration Date" (field 27) la ngay het hieu luc cua recurring journal, khong phai han dung.
    # Han dung cua lo nam o "Item Expiration Date" (field 6506), day moi la cot di vao
    # Item Ledger Entry."Expiration Date", xem ItemJnlPostLine.Codeunit.al dong 2021.
    head, tail, r1, r3, hdr, data = parse_sheet(load("xl/worksheets/sheet5.xml"))
    jmap = {
        "Posting Date": "Posting Date", "Document Date": "Posting Date",
        "Entry Type": "Entry Type", "Document No.": "Document No.",
        "Item No.": "Item No.", "Description": "Description",
        "Variant Code": "Variant Code", "Location Code": "Location Code",
        "Bin Code": "Bin Code", "Quantity": "Quantity",
        "Unit of Measure Code": "Unit of Measure Code", "Unit Amount": "Unit Amount",
        "Amount": "Amount", "Discount Amount": "Discount Amount", "Unit Cost": "Unit Cost",
        "Serial No.": "Serial No.", "Lot No.": "Lot No.",
        "Item Expiration Date": "Expiration Date", "Warranty Date": "Warranty Date",
    }
    for name in jmap:
        assert name in hdr, "sheet Item Journal Line khong co cot " + name
    rows = []
    for i, j in enumerate(journals, start=1):
        r = {hdr["Journal Template Name"]: TEMPLATE,
             hdr["Journal Batch Name"]: BATCH,
             hdr["Line No."]: str(i * 10000)}
        for name, key in jmap.items():
            v = fmt(j.get(key))
            if v != "":
                r[hdr[name]] = v
        rows.append(r)
    parts["xl/worksheets/sheet5.xml"] = build_sheet(head, tail, r1, r3, rows).encode("utf-8")
    parts["xl/tables/table5.xml"] = retable(load("xl/tables/table5.xml"), 3 + len(rows)).encode("utf-8")
    report["83 Item Journal Line"] = len(rows)

    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in order:
            zout.writestr(n, parts[n])

    print("\nket qua:")
    for k, v in report.items():
        print("  %-26s %6d dong" % (k, v))
    print("\nfile: %s (%.1f MB)" % (OUT, OUT.stat().st_size / 1048576))


if __name__ == "__main__":
    main()
