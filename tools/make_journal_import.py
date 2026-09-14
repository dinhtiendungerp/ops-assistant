"""Sinh file Table data for Item Journal Line de import lai bo du lieu demo cho cac ma co lot.

Dung khi phai post lai mot phan bo du lieu: xoa Item Ledger Entry cua nhung ma can lam lai, gan
Item Tracking Code, roi import file nay vao batch NWVDEMO va post. Cac ma khong nam trong file
giu nguyen ILE dang co.

Hai nhom dong trong file:
  - ma trong `TRACKED`: co Lot No. va Expiration Date.
  - ma truyen qua `--also`: de trong hai cot do. Dung cho ma da bi xoa ILE nhung khong quan ly
    lo, vi du 10000 va 10045 sau khi Dung bo chung khoi danh sach tracking ngay 12/09/2026.
    Phai de trong that: `ItemLedgEntry.CopyTrackingFromItemJnlLine` van chep Lot No. tu dong
    journal sang ILE ke ca khi item khong co tracking code, tuc se ra ILE co lo ma Item Card
    thi khong quan ly lo.

Hai cho phai dung, sai mot trong hai la mat du lieu ma khong co gi bao loi:

1. Cot han dung la **Expiration Date** (field 44), dung nhu ban xuat cua BC. Tooltip cua field
   do noi ve recurring journal, nhung khi batch bat "Item Tracking on Lines" thi no chinh la o
   nhap han dung cua lo. Chuoi day du:
     - `ItemJnlPostBatch` dong 975: batch bat co thi goi `ItemJournalLine.CreateItemTrackingLines`.
     - `ItemJnlLineReserve.CreateItemTracking` dong 446:
       `TempTrackingSpecification."Expiration Date" := ItemJournalLine."Expiration Date"`, tuc
       doc field 44 roi tao Reservation Entry.
     - `ItemJnlPostLine` dong 5814: `TempSplitItemJnlLine."Item Expiration Date" :=
       TempTrackingSpecification."Expiration Date"`, tuc 6506 duoc suy ra chu khong phai nhap tay.
     - `ItemJnlPostLine` dong 2021: `ItemLedgEntry."Expiration Date" :=
       ItemJnlLine."Item Expiration Date"`.
   Field 6506 `Editable = false`, ghi thang vao no la di duong tat: khong co Reservation Entry
   nen khong co Lot No. Information, va gia tri vua ghi bi buoc tao tracking ghi de len.
   Cung o `CreateItemTracking` dong 437: item khong co Item Tracking Code thi ham thoat ngay,
   lot tren dong bi bo qua ma khong bao loi. Vi vay phai gan tracking code truoc khi post.

2. Batch NWVDEMO phai bat **Item Tracking on Lines**, neu khong thi OnValidate cua field 6501
   xoa sach Lot No. ngay luc chen dong. App `NWV Marou Demo Setup` 1.0.1.0 bat san co nay.

Thu tu dong: giu nguyen thu tu cua bay file nguon, tuc tang dan theo ngay va trong mot ngay thi
nhap truoc xuat sau. Sap xep lai la co nguy co am ton giua chung.

Chay:
    python tools/make_journal_import.py                    # ra demo-data-nwv/import-lot-journal.txt
    python tools/make_journal_import.py --out <duong dan>
    python tools/make_journal_import.py --template <file xuat tu BC>   # doi chieu ten cot
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from demo_scenario import BATCH, TRACKED  # noqa: E402

DEMO = ROOT / "demo-data-nwv"
TEMPLATE_NAME = "ITEM"
LINE_STEP = 10000

# Dung thu tu cot cua ban xuat tu BC, khong them bot cot nao.
COLUMNS = ["Journal Template Name", "Journal Batch Name", "Line No.", "Item No.", "Posting Date",
           "Entry Type", "Document No.", "Location Code", "Quantity", "Unit Amount", "Unit Cost",
           "Amount", "Discount Amount", "Expiration Date", "Lot No."]

INBOUND = {"Purchase", "Positive Adjmt."}


def us(d) -> str:
    if isinstance(d, datetime):
        d = d.date()
    if not isinstance(d, date):
        d = date.fromisoformat(str(d)[:10])
    return f"{d.month:02d}/{d.day:02d}/{d.year % 100:02d}"


def num(v) -> str:
    """So viet kieu Anh, dau cham thap phan, bo phan .0 thua cho gon."""
    if v in (None, ""):
        return "0"
    f = float(v)
    return str(int(f)) if f == int(f) else f"{f:g}"


ALL = object()   # dau hieu lay het moi mat hang, xem --all


def doc_rows(them=None):
    import openpyxl

    lay = None if them is ALL else TRACKED | (them or set())
    rows = []
    for f in sorted(DEMO.glob("0[1-7]-ItemJournals-*.xlsx")):
        wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        for r in it:
            if not r or r[0] is None:
                continue
            d = dict(zip(hdr, r))
            if lay is not None and str(d["Item No."]) not in lay:
                continue
            rows.append(d)
        wb.close()
    return rows


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(DEMO / "import-lot-journal.txt"))
    ap.add_argument("--template", help="file Table data xuat tu BC, de doi chieu ten cot")
    ap.add_argument("--also", default="", help="ma khong quan ly lo nhung van can post lai, cach nhau dau phay")
    ap.add_argument("--all", action="store_true",
                    help="lay het moi mat hang, dung khi xoa sach ILE va post lai tu dau")
    args = ap.parse_args()

    them = ALL if args.all else {m.strip() for m in args.also.split(",") if m.strip()}
    if them is not ALL:
        trung = them & TRACKED
        if trung:
            sys.exit(f"--also trung voi TRACKED: {sorted(trung)}")
    src = doc_rows(them)
    n_lot = sum(1 for d in src if str(d["Item No."]) in TRACKED)
    ma_khong_lot = {str(d["Item No."]) for d in src} - TRACKED
    print(f"Doc {len(src)} dong: {n_lot} dong co lot ({len(TRACKED)} ma), "
          f"{len(src) - n_lot} dong khong lot ({len(ma_khong_lot)} ma)")

    thieu = [d for d in src if str(d["Item No."]) in TRACKED
             and (not d.get("Lot No.") or not d.get("Expiration Date"))]
    if thieu:
        sys.exit(f"{len(thieu)} dong cua ma co tracking bi thieu Lot No. hoac Expiration Date")

    dates = [d["Posting Date"] for d in src]
    assert all(dates[i] <= dates[i + 1] for i in range(len(dates) - 1)), "Posting Date khong tang dan"

    if args.template:
        head = Path(args.template).read_text(encoding="utf-8-sig").splitlines()[1]
        cot_bc = [c.lstrip("*") for c in head.split("\t")]
        khac = [c for c in cot_bc if c not in COLUMNS]
        if khac:
            print(f"Cot co trong file BC ma khong co o day: {khac}")
        thieu_cot = [c for c in COLUMNS if c not in cot_bc]
        if thieu_cot:
            print(f"Cot can them vao ban xuat tu BC: {thieu_cot}")

    out_lines = ["TABLEID:Item Journal Line", "\t".join("*" + c for c in COLUMNS)]
    ton = Counter()
    ton_lo = Counter()      # theo tung lo: BC ap dong xuat vao dong nhap cua DUNG lo do
    am = am_lo = 0
    for i, d in enumerate(src, start=1):
        qty = float(d["Quantity"])
        key = (str(d["Item No."]), d.get("Location Code") or "")
        vao = d["Entry Type"] in INBOUND
        ton[key] += qty if vao else -qty
        if ton[key] < -0.0001:
            am += 1
        if d.get("Lot No.") and str(d["Item No."]) in TRACKED:
            k2 = key + (str(d["Lot No."]),)
            ton_lo[k2] += qty if vao else -qty
            if ton_lo[k2] < -0.0001:
                am_lo += 1
        out_lines.append("\t".join([
            TEMPLATE_NAME, BATCH, str(i * LINE_STEP), str(d["Item No."]), us(d["Posting Date"]),
            str(d["Entry Type"]), str(d.get("Document No.") or ""), str(d.get("Location Code") or ""),
            num(qty), num(d.get("Unit Amount")), num(d.get("Unit Cost")), num(d.get("Amount")),
            num(d.get("Discount Amount")),
            us(d["Expiration Date"]) if str(d["Item No."]) in TRACKED else "",
            str(d["Lot No."]) if str(d["Item No."]) in TRACKED else "",
        ]))
    if am:
        sys.exit(f"{am} dong lam am ton Item x Location, thu tu dong sai, dung lai")
    if am_lo:
        # `ItemJnlPostLine` dong 2081 bao "cannot be fully applied" khi mot dong xuat cua lo
        # khong tim du so luong dang mo cua dung lo do tai dung dia diem.
        sys.exit(f"{am_lo} dong lam am ton theo LO, post se bao cannot be fully applied, dung lai")

    out = Path(args.out)
    out.write_text("\n".join(out_lines) + "\n", encoding="utf-8")
    print(f"Ton khong am o ca {len(ton)} cap Item x Location va {len(ton_lo)} to hop lo x dia diem")
    print(f"Line No. tu {LINE_STEP} den {len(src) * LINE_STEP}")
    print(f"\nDa ghi {len(src)} dong: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
