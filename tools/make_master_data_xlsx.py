"""Sinh file Excel master data cho bo du lieu demo Marou: nhom hang, dia diem, mat hang,
don vi tinh va Stockkeeping Unit. Import truoc, roi moi import Item Journal.

Nguong bo sung hang tren Stockkeeping Unit tinh tu chinh luong ban 90 ngay gan nhat trong
kich ban, khong phai so bia: reorder point = ban binh quan ngay x (lead time 5 + an toan 3),
reorder quantity = ban binh quan ngay x 14.
"""
from __future__ import annotations

import argparse
import random
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python"))
from bc_agent.make_fixtures import ITEMS, SEED, STORES, TODAY, WH, gen_sales  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

CATEGORIES = [
    ("BAR", "Thanh chocolate"),
    ("CONF", "Bonbon va keo"),
    ("DRINK", "Do uong"),
    ("GIFT", "Hop qua"),
    ("INGR", "Nguyen lieu"),
]

LOCATIONS = [
    ("WH-HCM", "Kho trung tam HCM", "Quan 7", "Ho Chi Minh"),
    ("S-CALMETTE", "Cua hang Calmette", "169 Calmette", "Ho Chi Minh"),
    ("S-THAODIEN", "Cua hang Thao Dien", "90 Xuan Thuy", "Ho Chi Minh"),
    ("S-HANOI", "Cua hang Ha Noi", "91 Tho Nhuom", "Ha Noi"),
    ("S-DANANG", "Cua hang Da Nang", "18 Bach Dang", "Da Nang"),
]

TRACKING_CODE = "LOTEXP"
GEN_PROD_GROUP = "RETAIL"
INV_POSTING_GROUP = "RESALE"
LEAD_DAYS, SAFETY_DAYS, TARGET_DAYS = 5, 3, 14
HIST_DAYS = 90
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
TABS = {"Item Category": "7F7F7F", "Location": "7F7F7F", "Item": "C00000",
        "Item Unit of Measure": "7F7F7F", "Stockkeeping Unit": "C00000"}


def sheet(wb, title: str, columns: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = TABS.get(title, "7F7F7F")
    ws.append(columns)
    for c in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for c, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c)].width = 26 if len(name) > 16 else 16
    ws.freeze_panes = "A2"


def avg_daily() -> dict[tuple[str, str], float]:
    rng = random.Random(SEED)
    sales = gen_sales(rng, days=180)
    since = TODAY - timedelta(days=HIST_DAYS)
    tot: dict[tuple[str, str], int] = defaultdict(int)
    for r in sales:
        if date.fromisoformat(r["date"]) >= since:
            tot[(r["item_no"], r["location_code"])] += r["qty"]
    return {k: v / HIST_DAYS for k, v in tot.items()}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheet(wb, "Item Category", ["Code", "Description"], [list(c) for c in CATEGORIES])
    sheet(wb, "Location", ["Code", "Name", "Address", "City"], [list(l) for l in LOCATIONS])

    item_cols = ["No.", "Description", "Base Unit of Measure", "Item Category Code", "Type",
                 "Costing Method", "Unit Cost", "Last Direct Cost", "Unit Price",
                 "Gen. Prod. Posting Group", "Inventory Posting Group",
                 "Item Tracking Code", "Expiration Calculation",
                 "Replenishment System", "Reordering Policy"]
    item_rows = []
    for item_no, desc, cat, cost, _, shelf in ITEMS:
        item_rows.append([item_no, desc, "PCS", cat, "Inventory", "FIFO", cost, cost,
                          round(cost * 2.2, 0), GEN_PROD_GROUP, INV_POSTING_GROUP,
                          TRACKING_CODE, f"<{shelf}D>", "Purchase", "Fixed Reorder Qty."])
    sheet(wb, "Item", item_cols, item_rows)

    sheet(wb, "Item Unit of Measure", ["Item No.", "Code", "Qty. per Unit of Measure"],
          [[i[0], "PCS", 1] for i in ITEMS])

    avg = avg_daily()
    sku_cols = ["Location Code", "Item No.", "Variant Code", "Replenishment System",
                "Reordering Policy", "Reorder Point", "Reorder Quantity",
                "Safety Stock Quantity", "Maximum Inventory", "Transfer-from Code"]
    sku_rows = []
    for loc in [WH] + STORES:
        for item_no, _, _, _, _, _ in ITEMS:
            a = avg.get((item_no, loc), 0.0)
            rp = round(a * (LEAD_DAYS + SAFETY_DAYS))
            rq = round(a * TARGET_DAYS)
            ss = round(a * SAFETY_DAYS)
            repl = "Purchase" if loc == WH else "Transfer"
            sku_rows.append([loc, item_no, "", repl, "Fixed Reorder Qty.",
                             rp, max(rq, 1), ss, round(a * 45), "" if loc == WH else WH])
    sheet(wb, "Stockkeeping Unit", sku_cols, sku_rows)

    path = out / "00-MasterData.xlsx"
    wb.save(path)
    print(f"{path}  |  {len(item_rows)} item, {len(sku_rows)} SKU, {len(LOCATIONS)} location")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
