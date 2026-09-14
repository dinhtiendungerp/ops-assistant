"""Doc bo du lieu demo trong `demo-data-nwv/` (bay file Item Journal va master data).

Dung o hai cho:
  - `tools/uc2_offline.py`: tinh ket qua mong doi de doi chieu voi BC sau khi import.
  - `bc_agent/make_fixtures.py`: dung mock cho tro ly tu chinh bo du lieu nay, de mock va BC that
    ra cung mot con so. Hom import xong chi doi BC_MODE, khong doi gi khac.

Moi dong Item Journal sau khi post thanh mot dong Item Ledger Entry, nen doi journal sang
hinh dang ILE (Quantity mang dau theo Entry Type) la du de chay lop tinh toan.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DEMO_DIR = ROOT / "demo-data-nwv"
INBOUND = {"Purchase", "Positive Adjmt."}


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v:
        return date.fromisoformat(str(v)[:10])
    return None


def read_items(demo_dir: Path = DEMO_DIR) -> dict[str, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(demo_dir / "00-MasterData.xlsx", data_only=True)
    rows = list(wb["Item"].iter_rows(values_only=True))
    hdr = list(rows[0])
    out = {}
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        if not d.get("No."):
            continue
        out[str(d["No."])] = {
            "item_no": str(d["No."]),
            "description": d.get("Description") or "",
            "unit_cost": float(d.get("Unit Cost") or 0),
            "unit_price": float(d.get("Unit Price") or 0),
            "item_tracking_code": d.get("Item Tracking Code") or "",
        }
    return out


def read_ile(demo_dir: Path = DEMO_DIR) -> list[dict]:
    """Doi dong Item Journal thanh dong Item Ledger Entry: Quantity mang dau theo Entry Type."""
    import openpyxl

    out: list[dict] = []
    for f in sorted(demo_dir.glob("0[1-7]-ItemJournals-*.xlsx")):
        wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        for r in it:
            if not r or r[0] is None:
                continue
            d = dict(zip(hdr, r))
            qty = float(d["Quantity"])
            out.append({
                "posting_date": as_date(d["Posting Date"]),
                "entry_type": d["Entry Type"],
                "item_no": str(d["Item No."]),
                "location_code": d.get("Location Code") or "",
                "lot_no": d.get("Lot No.") or None,
                "expiration_date": as_date(d.get("Expiration Date")),
                "quantity": qty if d["Entry Type"] in INBOUND else -qty,
                "item_category": "",
            })
        wb.close()
    return out
